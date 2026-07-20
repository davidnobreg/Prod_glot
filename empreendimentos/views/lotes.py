import json
import re
from datetime import date
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib import messages
import pandas as pd
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.db.models import Q
from django.views.decorators.http import require_http_methods
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation

from rolepermissions.decorators import has_permission_decorator

from django.core.paginator import Paginator, PageNotAnInteger, EmptyPage
from django.views import View

from django.db import transaction
from django.utils import timezone
from datetime import datetime, timedelta

from ..forms import ArquivoForm, LoteForm, AtualizarLoteForm
from ..models import Empreendimento, Quadra, Lote, TypeLote
from vendas.models import RegisterVenda


@has_permission_decorator('listaQuadra')
def listaQuadra(request, empreendimento_uuid):
    empreendimento = get_object_or_404(Empreendimento, uuid=empreendimento_uuid)
    situacao_filtro = request.GET.get('situacao')

    def formatar_moeda(valor):
        return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    # Inicializa a consulta de quadras
    quadras = Quadra.objects.filter(empr=empreendimento).order_by('id')

    quadras_info_list = []

    for quadra in quadras:
        # Filtra os lotes diretamente aqui, com base na situação, se existir
        lotes = quadra.lotes.all()

        if situacao_filtro and situacao_filtro != 'TODOS':
            if situacao_filtro == 'OUTROS':
                lotes = lotes.filter(
                    Q(situacao='CONSTRUTORA') |
                    Q(situacao='EM_RESERVA') |
                    Q(situacao='INDISPONIVEL')
                )
            else:
                lotes = lotes.filter(situacao=situacao_filtro)

        # Se não houver lotes após o filtro, pula para a próxima quadra
        if not lotes.exists():
            continue

        total_livres = lotes.filter(situacao="DISPONIVEL").count()
        total_vendidos = lotes.filter(situacao="VENDIDO").count()
        outros = lotes.filter(
            Q(situacao='CONSTRUTORA') |
            Q(situacao='EM_RESERVA') |
            Q(situacao='INDISPONIVEL')
        ).count()

        lotes_info = []

        for lote in lotes:
            try:
                area = Decimal(str(lote.area or "0").replace(",", "."))
                valor_metro = Decimal(str(lote.valor_metro_quadrado or "0").replace(",", "."))

                quantidade_parcela = lote.quadra.empr.quantidade_parcela or 0
                quantidade_parcela = int(quantidade_parcela)

                valor_total = area * valor_metro

                if quantidade_parcela > 0:
                    valor_parcela = (valor_total / Decimal(quantidade_parcela)).quantize(
                        Decimal("0.01"),
                        rounding=ROUND_HALF_UP
                    )
                else:
                    valor_parcela = Decimal("0.00")

            except (TypeError, ValueError, InvalidOperation, AttributeError):
                valor_parcela = Decimal("0.00")

            lotes_info.append({
                "lote": lote,
                "situacao": lote.situacao,
                "valor_parcela_formatado": formatar_moeda(valor_parcela),
            })



        quadras_info_list.append({
            'quadra': quadra,
            'total_livres': total_livres,
            'total_vendidos': total_vendidos,
            'lotes': lotes_info,
            'outros': outros
        })

    # Paginação
    paginator = Paginator(quadras_info_list, 9)
    page = request.GET.get('page')

    try:
        quadras_info_page = paginator.page(page)
    except PageNotAnInteger:
        quadras_info_page = paginator.page(1)
    except EmptyPage:
        quadras_info_page = paginator.page(paginator.num_pages)

    # Estatísticas gerais
    all_lotes = Lote.objects.filter(quadra__empr=empreendimento)

    querydict = request.GET.copy()  # Torna o QueryDict mutável

    if 'page' in querydict:
        querydict.pop('page')  # Remove o parâmetro 'page'

    context = {
        'quadras_info': quadras_info_page,
        'empreendimento': empreendimento,
        'total': all_lotes.count(),
        'livre': all_lotes.filter(situacao='DISPONIVEL').count(),
        'prereserva': all_lotes.filter(situacao='PRE-RESERVA').count(),
        'reservado': all_lotes.filter(situacao='RESERVADO').count(),
        'pre_venda': all_lotes.filter(situacao='PRE-VENDA').count(),
        'vendido': all_lotes.filter(situacao='VENDIDO').count(),
        'analise': all_lotes.filter(situacao='ANALISE').count(),
        'outros': all_lotes.filter(
            Q(situacao='CONSTRUTORA') |
            Q(situacao='INDISPONIVEL')
        ).count(),
        'situacao_filtro': situacao_filtro,
        'querystring': querydict.urlencode()
    }

    return render(request, 'lista-quadras.html', context)


@has_permission_decorator('atualizarLotes')
def atualizarLotes(request, empreendimento_uuid):
    empreendimento = get_object_or_404(Empreendimento, uuid=empreendimento_uuid)
    quadra = request.GET.get('quadra', '').strip()
    lote = request.GET.get('lote', '').strip()

    lotes = Lote.objects.select_related('quadra', 'quadra__empr').filter(
        quadra__empr=empreendimento
    ).order_by(
        'id',
        'quadra__namequadra',
        'lote'
    )

    if quadra:
        lotes = lotes.filter(quadra__namequadra__iexact=quadra)

    if lote:
        lotes = lotes.filter(lote__iexact=lote)

    paginator = Paginator(lotes, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    querydict = request.GET.copy()
    querydict.pop('page', None)

    context = {
        'lotes': page_obj,
        'page_obj': page_obj,
        'empreendimento': empreendimento,
        'filtro_quadra': quadra,
        'filtro_lote': lote,
        'querystring': querydict.urlencode(),
    }
    return render(request, 'atualizar-lotes.html', context)


@has_permission_decorator('atualizarLotes')
def editarAtualizarLote(request, lote_uuid):
    lote = get_object_or_404(Lote.objects.select_related('quadra', 'quadra__empr'), uuid=lote_uuid)
    querystring = request.GET.urlencode()

    if request.method == 'POST':
        form = AtualizarLoteForm(request.POST, instance=lote)
        querystring = request.POST.get('querystring', '')

        if form.is_valid():
            form.save()
            messages.success(request, "Lote atualizado com sucesso!")
            url = reverse('atualizar-lotes', args=[lote.quadra.empr.uuid])
            if querystring:
                url = f'{url}?{querystring}'
            return redirect(url)

        messages.error(request, "Verifique os campos informados.")
    else:
        form = AtualizarLoteForm(instance=lote)

    context = {
        'form': form,
        'lote': lote,
        'querystring': querystring,
    }
    return render(request, 'editar-atualizar-lote.html', context)


class importarDados(View):
    template_name = 'empreendimento_arq.html'

    def get(self, request, uuid, *args, **kwargs):
        empreendimento = get_object_or_404(Empreendimento, uuid=uuid)
        form = ArquivoForm()
        context = {'form': form, 'empreendimento': empreendimento}
        return render(request, self.template_name, context)

    def post(self, request, uuid):
        empreendimento = get_object_or_404(Empreendimento, uuid=uuid)
        form = ArquivoForm(request.POST, request.FILES)
        if form.is_valid():
            arquivo = request.FILES['arquivo']

            # Lê o Excel a partir da terceira linha (pulando as duas primeiras)
            df = pd.read_excel(arquivo, skiprows=1,
                               names=["quadra", "lote", "area", "valor_metro_quadrado", "situacao"])

            # Remover linhas vazias
            df = df.dropna(subset=["quadra", "lote", "area", "valor_metro_quadrado"])

            # Converter os valores para string para evitar problemas
            df["quadra"] = df["quadra"].astype(str).str.strip()
            df["lote"] = df["lote"].astype(str).str.strip()

            # Iterar sobre as linhas do DataFrame e criar registros
            for _, row in df.iterrows():
                self.criar_quadra(row, empreendimento.id)
            messages.success(request, "Arquivo importado com sucesso!")
            return redirect('lista-empreendimento-tabela')  # Ajuste para onde quer redirecionar

        return render(request, self.template_name, {'form': form})

    def criar_quadra(self, row, empreendimento_id):
        quadra, _ = Quadra.objects.get_or_create(
            namequadra=row["quadra"],
            empr_id=empreendimento_id  # Supondo que 'empr' seja uma ForeignKey para 'empreendimento'
        )

        # Criar o lote com a situação original do arquivo
        Lote.objects.create(
            lote=row["lote"],
            area=row["area"],
            situacao=row["situacao"],  # Mantém o status real do lote
            valor_metro_quadrado=row["valor_metro_quadrado"],
            quadra=quadra
        )


@has_permission_decorator('reservarLote')
def alteraLote(request, uuid):
    lote = get_object_or_404(Lote, uuid=uuid)
    get_tempo = Empreendimento.objects.get(id=lote.quadra.empr_id)

    try:
        area = float(lote.area)
        valor_metro = float(lote.valor_metro_quadrado)
        valor = area * valor_metro
    except (TypeError, ValueError):
        valor = 0

    # Formatação para moeda brasileira
    valor_formatado = f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    # Pegando a quantidade de parcelas (ajuste o nome conforme seu modelo)
    try:
        total_parcelas = int(get_tempo.quantidade_parcela)
    except (TypeError, ValueError, AttributeError):
        total_parcelas = 0

    # Cálculo do valor da parcela
    if total_parcelas > 0:
        valor_parcela = valor / total_parcelas
    else:
        valor_parcela = 0

    # Formatação do valor da parcela
    valor_parcela_formatado = f"R$ {valor_parcela:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    # Caso especial: liberar lote se estiver em reserva mas sem lógica de uso (defensivo)
    if lote.situacao == "EM_RESERVA" and not lote.user:
        lote.situacao = "DISPONIVEL"
        lote.save()
        messages.error(request, "Pre-Reservado Cancelada!")

    if request.method == 'GET':
        # Ao acessar, define como EM_RESERVA
        if lote.situacao == "DISPONIVEL":
            lote.situacao = "EM_RESERVA"
            lote.tempo_reservado = timezone.now() + timedelta(days=lote.quadra.empr.tempo_reserva)
            lote.save()
        form = LoteForm(instance=lote)
        context = {'form': form,
                   'lote': lote,
                   'valor_formatado': valor_formatado,
                   'valor_parcela_formatado': valor_parcela_formatado,
                   'total_parcelas': total_parcelas}
        return render(request, 'reserva-temporaria.html', context)

    elif request.method == 'POST':
        form = LoteForm(request.POST, request.FILES, instance=lote)
        if form.is_valid():
            lote = form.save(commit=False)
            lote.situacao = 'PRE-RESERVA'
            lote.data_termina_reserva = timezone.now() + timedelta(days=lote.quadra.empr.tempo_reserva)
            lote.user = request.user.first_name
            lote.telefone_user = request.user.contato
            lote.save()
            messages.success(request, "Pre-Reservado Salva Com Sucesso!")
            return redirect('listar-quadras', uuid=lote.quadra.empr_uuid)

        else:
            context = {'form': form,
                       'lote': lote,
                       'valor_formatado': valor_formatado,
                       'valor_parcela_formatado': valor_parcela_formatado,
                       'total_parcelas': total_parcelas}
            return render(request, 'reserva-temporaria.html', context)


@has_permission_decorator('reservadoDetalheEmpreendimento')
def reservadoDetalheEmpreendimento(request, preReserva_uuid):
    #lote = Lote.objects.filter(uuid=preReserva_uuid).first()

    lote = None

    from django.http import HttpResponseForbidden

    if not request.user.is_authenticated:
        lote = Lote.objects.filter(uuid=preReserva_uuid).first()
        context = {'lote': lote}
        return render(request, 'permissao.html', context)
        #return HttpResponseForbidden("Você não tem permissão.")

    if request.user.tipo_usuario == "ADMINISTRADOR":
        lote = Lote.objects.filter(uuid=preReserva_uuid).first()
    else:
        lote = Lote.objects.filter(
            uuid=preReserva_uuid,
            user=request.user.first_name
        ).first()

    if not lote:
        lote = Lote.objects.filter(uuid=preReserva_uuid).first()

        context = {'lote': lote}
        return render(request, 'permissao.html', context)
        #return HttpResponseForbidden("Você não tem permissão para acessar este lote.")

    context = {'lote': lote}

    return render(request, 'detalhes-reserva-lote.html', context)


@has_permission_decorator('listaReservasTemporaria')
def listaReservasTemporaria(request):
    query = request.GET.get('q', '')
    tipo_empreendimento = request.GET.get('tipo_empreendimento', '')

    # Filtra lotes 'PRE-RESERVA'
    lotes = Lote.objects.filter(situacao='PRE-RESERVA')

    if query:
        lotes = lotes.filter(
            Q(quadra__empr__nome__icontains=query) |
            Q(quadra__namequadra__icontains=query) |
            Q(lote__icontains=query) |
            Q(user__icontains=query)
        )

    if tipo_empreendimento:
        lotes = lotes.filter(quadra__empr__id=tipo_empreendimento)

    paginator = Paginator(lotes, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    empreendimentos = Empreendimento.objects.filter(is_ativo=False).order_by('id')

    context = {
        'page_obj': page_obj,
        'query': query,
        'tipo_empreendimento': tipo_empreendimento,
        'empreendimentos': empreendimentos,
    }
    return render(request, 'relatorio_de_reservas_temporario.html', context)


@has_permission_decorator('liberaLote')
def liberaLote(request, lote_uuid):
    get_lote = get_object_or_404(Lote, uuid=lote_uuid)

    if request.method == 'GET':
        get_lote.situacao = "DISPONIVEL"
        get_lote.save()
        messages.success(request, "Lote liberado com Sucesso!")
    return redirect('listar-quadras', empreendimento_uuid=get_lote.quadra.empr.uuid)


@has_permission_decorator('cancelarReservadoTemporaria')
def cancelarReservadoTemporaria(request, lote_uuid):
    get_lote = get_object_or_404(Lote, uuid=lote_uuid)

    if request.method == 'GET':
        get_lote.situacao = "DISPONIVEL"
        get_lote.cliente_reserva = ""
        get_lote.telefone = ""
        get_lote.save()
        messages.error(request, "Pre-Resevado Cancelada!")
    return redirect('listar-quadras', empreendimento_uuid=get_lote.quadra.empr.uuid)


@has_permission_decorator('cancelarReservadoTemporariaLista')
def cancelarReservadoTemporariaLista(request, lote_uuid):
    get_lote = get_object_or_404(Lote, uuid=lote_uuid)

    if request.method == 'GET':
        get_lote.situacao = "DISPONIVEL"
        get_lote.cliente_reserva = ""
        get_lote.telefone = ""
        get_lote.save()
        messages.error(request, "Pre-Resevado Cancelada!")
    return redirect('lista-pre-reserva')


@has_permission_decorator('renovarReservaTemporaria')
def renovaReserva(request, renova_uuid):
    get_lote = Lote.objects.get(uuid=renova_uuid)

    get_lote.data_termina_reserva = datetime.now() + timedelta(days=get_lote.quadra.empr.tempo_reserva)

    get_lote.save()
    messages.success(request, "Reserva renovada com sucesso!")
    return redirect('lista-pre-reserva')

def gerarRelatorioLotes(request):

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = (
        'attachment; filename="relatorio_lotes.pdf"'
    )

    doc = SimpleDocTemplate(
        response,
        pagesize=A4,
        leftMargin=30,
        rightMargin=30,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    elementos = []

    # filtros
    situacao = request.GET.get('situacao', 'TODOS')
    loteamento_uuid = request.GET.get('loteamento_uuid')
    empreendimento = None

    if loteamento_uuid:
        empreendimento = Empreendimento.objects.filter(uuid=loteamento_uuid).first()

    # consulta inicial
    lotes = Lote.objects.select_related(
        'quadra',
        'quadra__empr'
    ).all()

    if empreendimento:
        lotes = lotes.filter(
            quadra__empr=empreendimento
        )

    # filtra situação se não for TODOS
    if situacao != 'TODOS':
        if situacao == 'OUTROS':
            lotes = lotes.filter(
                Q(situacao='CONSTRUTORA') |
                Q(situacao='EM_RESERVA') |
                Q(situacao='INDISPONIVEL')
            )
        else:
            lotes = lotes.filter(situacao=situacao)

    # nome empreendimento
    primeiro_lote = lotes.first()

    nome_empreendimento = (
        empreendimento.nome
        if empreendimento
        else primeiro_lote.quadra.empr.nome
        if primeiro_lote
        else 'Empreendimento nao identificado'
    )
    title_style = ParagraphStyle(
        'RelatorioLotesTitle',
        parent=styles['Title'],
        alignment=1,
        fontSize=18,
        leading=22,
        spaceAfter=12
    )

    subtitle_style = ParagraphStyle(
        'RelatorioLotesSubtitle',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        spaceAfter=12
    )

    if empreendimento and empreendimento.logo:
        try:
            logo_data = BytesIO(empreendimento.logo.read())
            logo = Image(logo_data)
            logo.drawHeight = 58
            logo.drawWidth = 150
            logo.hAlign = 'CENTER'
            elementos.append(logo)
            elementos.append(Spacer(1, 12))
        except Exception:
            pass

    titulo = Paragraph(
        f'Relatorio de Lotes - <b>{nome_empreendimento}</b>',
        title_style
    )

    elementos.append(titulo)
    elementos.append(Spacer(1, 10))

    subtitulo = Paragraph(
        f'Situacao dos Lotes: <b>{situacao}</b>',
        subtitle_style
    )

    elementos.append(subtitulo)
    elementos.append(Spacer(1, 10))

    dados = [[
        'Quadra',
        'Lote',
        'Situacao',
        'Valor do Lote'
    ]]

    def formatar_moeda(valor):
        return f'R$ {valor:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

    def calcular_valor_lote(lote):
        try:
            area = float(str(lote.area or 0).replace(',', '.'))
            valor_metro = float(str(lote.valor_metro_quadrado or 0).replace(',', '.'))
            return area * valor_metro
        except (TypeError, ValueError):
            return 0

    for lote in lotes:
        dados.append([
            lote.quadra.namequadra,
            lote.lote,
            lote.situacao,
            formatar_moeda(calcular_valor_lote(lote))
        ])

    tabela = Table(
        dados,
        colWidths=[105, 105, 150, 160],
        repeatRows=1
    )

    tabela.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#08789A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ALIGN', (0, 1), (2, -1), 'CENTER'),
        ('ALIGN', (3, 1), (3, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#9CA3AF')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 9),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))

    elementos.append(tabela)

    doc.build(elementos)

    return response


# ===================================================================
# Exportar / Importar Lotes em massa (xlsx)
# ===================================================================

_VENDA_ATIVA_TIPOS = {'ANALISE', 'PRE-VENDA'}
_STATUS_LOTE_VALIDOS = set(TypeLote.values)


def _lote_tem_venda_ativa(lote):
    return RegisterVenda.objects.filter(
        lote=lote,
        tipo_venda__in=_VENDA_ATIVA_TIPOS,
    ).exists()


@has_permission_decorator('atualizarLotes')
def exportar_lotes(request, empreendimento_uuid):
    empr = get_object_or_404(Empreendimento, uuid=empreendimento_uuid)
    lotes = (
        Lote.objects
        .filter(quadra__empr=empr)
        .select_related('quadra')
        .order_by('quadra__namequadra', 'lote')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Lotes'

    header_fill = PatternFill(fill_type='solid', fgColor='0F3460')
    header_font = Font(color='FFFFFF', bold=True)
    locked_fill = PatternFill(fill_type='solid', fgColor='D3D3D3')

    headers = ['id', 'numero', 'quadra', 'area', 'preco', 'status', 'medidas', 'confrontacoes', 'cliente_reserva', 'corretor',]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for row_idx, lote in enumerate(lotes, 2):
        bloqueado = _lote_tem_venda_ativa(lote)
        ws.cell(row=row_idx, column=1, value=lote.id)
        ws.cell(row=row_idx, column=2, value=lote.lote)
        ws.cell(row=row_idx, column=3, value=lote.quadra.namequadra)
        ws.cell(row=row_idx, column=4, value=lote.area)
        ws.cell(row=row_idx, column=5, value=lote.valor_metro_quadrado)
        status_cell = ws.cell(
            row=row_idx, column=6,
            value='[BLOQUEADO]' if bloqueado else lote.situacao,
        )
        if bloqueado:
            status_cell.fill = locked_fill
        ws.cell(row=row_idx, column=7, value=lote.medidas or '')
        ws.cell(row=row_idx, column=8, value=lote.confrontacoes or '')
        ws.cell(row=row_idx, column=9, value=lote.cliente_reserva or '')
        ws.cell(row=row_idx, column=10, value=lote.user or '')

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    date_str = date.today().strftime('%Y-%m-%d')
    nome_safe = re.sub(r'[^\w\s-]', '', empr.nome).strip().replace(' ', '_')
    filename = f'lotes_{nome_safe}_{date_str}.xlsx'

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@has_permission_decorator('atualizarLotes')
@require_http_methods(['POST'])
def importar_lotes(request, empreendimento_uuid):
    empr = get_object_or_404(Empreendimento, uuid=empreendimento_uuid)
    arquivo = request.FILES.get('arquivo')
    if not arquivo:
        messages.error(request, 'Selecione um arquivo xlsx.')
        return redirect(reverse('detalhe-empreendimento', args=[empr.uuid]))

    try:
        wb = openpyxl.load_workbook(arquivo, data_only=True)
        ws = wb.active
    except Exception:
        messages.error(request, 'Arquivo inválido. Envie um xlsx gerado pela exportação.')
        return redirect(reverse('detalhe-empreendimento', args=[empr.uuid]))

    ids_empr = set(
        Lote.objects.filter(quadra__empr=empr).values_list('id', flat=True)
    )

    alteracoes = []
    ignorados = []
    erros = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(cell is not None for cell in row):
            continue

        row_padded = (list(row) + [None] * 10)[:10]
        lote_id_raw, numero, quadra_nome, area, preco, status, medidas, confrontacoes, cliente_reserva_val, corretor = row_padded

        if lote_id_raw is None:
            erros.append({'linha': row_idx, 'motivo': 'ID ausente'})
            continue

        try:
            lote_id = int(lote_id_raw)
        except (TypeError, ValueError):
            erros.append({'linha': row_idx, 'motivo': f'ID inválido: {lote_id_raw}'})
            continue

        if lote_id not in ids_empr:
            erros.append({
                'linha': row_idx,
                'motivo': f'Lote id={lote_id} não pertence a este empreendimento',
            })
            continue

        try:
            lote = Lote.objects.select_related('quadra').get(id=lote_id)
        except Lote.DoesNotExist:
            erros.append({'linha': row_idx, 'motivo': f'Lote id={lote_id} não encontrado'})
            continue

        if _lote_tem_venda_ativa(lote):
            ignorados.append({'id': lote_id, 'numero': lote.lote, 'motivo': 'venda ativa'})
            continue

        campos = {}

        if area is not None:
            area_str = str(area).strip()
            if area_str != str(lote.area or '').strip():
                campos['area'] = {'atual': lote.area or '', 'novo': area_str}

        if preco is not None:
            preco_str = str(preco).strip()
            try:
                preco_val = Decimal(preco_str.replace(',', '.'))
            except InvalidOperation:
                erros.append({'linha': row_idx, 'motivo': f'Preço inválido: {preco_str}'})
                continue
            if preco_val < 0:
                erros.append({'linha': row_idx, 'motivo': f'Preço negativo: {preco_str}'})
                continue
            if preco_str != str(lote.valor_metro_quadrado or '').strip():
                campos['valor_metro_quadrado'] = {
                    'atual': str(lote.valor_metro_quadrado or ''),
                    'novo': preco_str,
                }

        if status is not None:
            status_str = str(status).strip()
            if status_str not in ('[BLOQUEADO]', ''):
                if status_str not in _STATUS_LOTE_VALIDOS:
                    erros.append({
                        'linha': row_idx,
                        'motivo': (
                            f'Status inválido: "{status_str}". '
                            f'Válidos: {", ".join(sorted(_STATUS_LOTE_VALIDOS))}'
                        ),
                    })
                    continue
                if status_str != lote.situacao:
                    campos['situacao'] = {'atual': lote.situacao, 'novo': status_str}

        if medidas is not None:
            medidas_str = str(medidas).strip()
            atual_medidas = (lote.medidas or '').strip()
            if medidas_str != atual_medidas:
                campos['medidas'] = {'atual': atual_medidas, 'novo': medidas_str}

        if confrontacoes is not None:
            confrontacoes_str = str(confrontacoes).strip()
            atual = (lote.confrontacoes or '').strip()
            if confrontacoes_str != atual:
                campos['confrontacoes'] = {'atual': atual, 'novo': confrontacoes_str}

        if cliente_reserva_val is not None:
            cr_str = str(cliente_reserva_val).strip()
            atual_cr = (lote.cliente_reserva or '').strip()
            if cr_str != atual_cr:
                campos['cliente_reserva'] = {'atual': atual_cr, 'novo': cr_str}

        if corretor is not None:
            corretor_str = str(corretor).strip()
            atual_cr_user = (lote.user or '').strip()
            if corretor_str != atual_cr_user:
                campos['user'] = {'atual': atual_cr_user, 'novo': corretor_str}

        if campos:
            alteracoes.append({'id': lote_id, 'numero': lote.lote, 'campos': campos})

    context = {
        'empreendimento': empr,
        'alteracoes': alteracoes,
        'ignorados': ignorados,
        'erros': erros,
        'alteracoes_json': json.dumps(alteracoes),
    }
    return render(request, 'importar-lotes-preview.html', context)


@has_permission_decorator('atualizarLotes')
@require_POST
def importar_lotes_confirmar(request, empreendimento_uuid):
    empr = get_object_or_404(Empreendimento, uuid=empreendimento_uuid)
    try:
        alteracoes = json.loads(request.POST.get('alteracoes_json', '[]'))
    except json.JSONDecodeError:
        messages.error(request, 'Dados inválidos. Refaça a importação.')
        return redirect(reverse('detalhe-empreendimento', args=[empr.uuid]))

    ids_empr = set(
        Lote.objects.filter(quadra__empr=empr).values_list('id', flat=True)
    )
    campos_permitidos = {'area', 'valor_metro_quadrado', 'situacao', 'medidas', 'confrontacoes', 'cliente_reserva', 'user'}

    atualizados = 0
    ignorados = 0
    with transaction.atomic():
        for alt in alteracoes:
            try:
                lote_id = int(alt['id'])
            except (KeyError, TypeError, ValueError):
                continue
            if lote_id not in ids_empr:
                continue
            try:
                lote = Lote.objects.get(id=lote_id)
            except Lote.DoesNotExist:
                continue
            if _lote_tem_venda_ativa(lote):
                ignorados += 1
                continue

            campos = alt.get('campos', {})
            update_fields = []
            for field, valores in campos.items():
                if field in campos_permitidos:
                    setattr(lote, field, valores.get('novo', ''))
                    update_fields.append(field)

            if update_fields:
                lote.save(update_fields=update_fields)
                atualizados += 1

    partes = [f'✅ {atualizados} lotes atualizados']
    if ignorados:
        partes.append(f'⚠️ {ignorados} ignorados (venda ativa)')
    messages.success(request, ' — '.join(partes))
    return redirect(reverse('detalhe-empreendimento', args=[empr.uuid]))
