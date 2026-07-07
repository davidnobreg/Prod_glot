import json
from decimal import Decimal
from io import BytesIO

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse

from empreendimentos.models import Empreendimento, Quadra, Lote, TypeLote
from vendas.models import RegisterVenda

User = get_user_model()

_STATUS_VALIDOS = set(TypeLote.values)


def _make_user(username='testuser'):
	return User.objects.create_user(
		username=username,
		password='pass123',
		email=f'{username}@test.com',
		is_superuser=True,
		is_staff=True,
	)


def _make_empreendimento(**kwargs):
	defaults = {
		'nome': 'Residencial Teste',
		'telefone': '(83) 99999-9999',
		'tempo_reserva': 30,
		'quantidade_parcela': 60,
	}
	defaults.update(kwargs)
	return Empreendimento.objects.create(**defaults)


def _make_quadra(empr, nome='Quadra A'):
	return Quadra.objects.create(namequadra=nome, empr=empr)


def _make_lote(quadra, numero='Lote 01', **kwargs):
	defaults = {
		'lote': numero,
		'area': '200',
		'situacao': TypeLote.DISPONIVEL,
		'valor_metro_quadrado': '1500.00',
		'telefone': '(83) 99999-9999',
		'telefone_user': '(83) 99999-9999',
		'user': 'teste',
		'cliente_reserva': '0',
	}
	defaults.update(kwargs)
	return Lote.objects.create(quadra=quadra, **defaults)


def _make_xlsx(rows, headers=None):
	"""Cria um xlsx em memória com rows (list of lists). headers opcional."""
	wb = openpyxl.Workbook()
	ws = wb.active
	if headers is None:
		headers = ['id', 'numero', 'quadra', 'area', 'preco', 'status', 'descricao']
	ws.append(headers)
	for row in rows:
		ws.append(row)
	buf = BytesIO()
	wb.save(buf)
	buf.seek(0)
	return buf


class ExportarLotesTest(TestCase):

	def setUp(self):
		self.user = _make_user()
		self.client.force_login(self.user)
		self.empr = _make_empreendimento()
		self.quadra = _make_quadra(self.empr)
		self.lote = _make_lote(self.quadra)
		self.url = reverse('exportar-lotes', args=[self.empr.uuid])

	def test_exportar_lotes_xlsx(self):
		response = self.client.get(self.url)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			response['Content-Type'],
			'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		)
		wb = openpyxl.load_workbook(BytesIO(response.content))
		ws = wb.active
		headers = [ws.cell(1, c).value for c in range(1, 8)]
		self.assertIn('id', headers)
		self.assertIn('status', headers)
		data_row = [ws.cell(2, c).value for c in range(1, 8)]
		self.assertEqual(data_row[0], self.lote.id)

	def test_exportar_lotes_bloqueado_com_venda_ativa(self):
		cliente = User.objects.first()
		RegisterVenda.objects.create(
			lote=self.lote,
			tipo_venda='ANALISE',
		)
		response = self.client.get(self.url)
		wb = openpyxl.load_workbook(BytesIO(response.content))
		ws = wb.active
		status_val = ws.cell(2, 6).value
		self.assertEqual(status_val, '[BLOQUEADO]')

	def test_anonimo_bloqueado(self):
		self.client.logout()
		response = self.client.get(self.url)
		self.assertEqual(response.status_code, 403)


@override_settings(
	DEFAULT_FILE_STORAGE='django.core.files.storage.FileSystemStorage',
	STORAGES={
		'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
		'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
	},
)
class ImportarLotesTest(TestCase):

	def setUp(self):
		self.user = _make_user()
		self.client.force_login(self.user)
		self.empr = _make_empreendimento()
		self.quadra = _make_quadra(self.empr)
		self.lote = _make_lote(self.quadra, area='200', valor_metro_quadrado='1500.00')
		self.url = reverse('importar-lotes', args=[self.empr.uuid])

	def _post_xlsx(self, rows):
		buf = _make_xlsx(rows)
		return self.client.post(self.url, {'arquivo': buf}, format='multipart')

	def test_importar_lotes_atualiza_dados(self):
		rows = [[self.lote.id, self.lote.lote, 'Quadra A', '250', '2000.00', 'DISPONIVEL', 'novo obs']]
		response = self._post_xlsx(rows)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, str(self.lote.id))
		self.assertContains(response, '250')
		self.assertContains(response, '2000.00')

	def test_importar_lotes_ignora_venda_ativa(self):
		RegisterVenda.objects.create(lote=self.lote, tipo_venda='ANALISE')
		rows = [[self.lote.id, self.lote.lote, 'Quadra A', '999', '9999.00', 'DISPONIVEL', '']]
		response = self._post_xlsx(rows)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'venda ativa')
		self.lote.refresh_from_db()
		self.assertEqual(self.lote.area, '200')

	def test_importar_lotes_erro_preco_invalido(self):
		rows = [[self.lote.id, self.lote.lote, 'Quadra A', '200', 'abc', 'DISPONIVEL', '']]
		response = self._post_xlsx(rows)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Preço inválido')

	def test_importar_lotes_erro_status_invalido(self):
		rows = [[self.lote.id, self.lote.lote, 'Quadra A', '200', '1500', 'INEXISTENTE', '']]
		response = self._post_xlsx(rows)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Status inválido')


@override_settings(
	DEFAULT_FILE_STORAGE='django.core.files.storage.FileSystemStorage',
	STORAGES={
		'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
		'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
	},
)
class ImportarLotesConfirmarTest(TestCase):

	def setUp(self):
		self.user = _make_user()
		self.client.force_login(self.user)
		self.empr = _make_empreendimento()
		self.quadra = _make_quadra(self.empr)
		self.lote = _make_lote(self.quadra, area='200', valor_metro_quadrado='1500.00')
		self.url = reverse('importar-lotes-confirmar', args=[self.empr.uuid])

	def _post_confirmar(self, alteracoes):
		return self.client.post(self.url, {
			'alteracoes_json': json.dumps(alteracoes),
		})

	def test_confirmar_aplica_alteracoes(self):
		alteracoes = [{
			'id': self.lote.id,
			'numero': self.lote.lote,
			'campos': {
				'area': {'atual': '200', 'novo': '350'},
				'valor_metro_quadrado': {'atual': '1500.00', 'novo': '2500.00'},
			},
		}]
		response = self._post_confirmar(alteracoes)
		self.assertRedirects(
			response,
			reverse('detalhe-empreendimento', args=[self.empr.uuid]),
			fetch_redirect_response=False,
		)
		self.lote.refresh_from_db()
		self.assertEqual(self.lote.area, '350')
		self.assertEqual(self.lote.valor_metro_quadrado, '2500.00')

	def test_confirmar_ignora_venda_ativa(self):
		RegisterVenda.objects.create(lote=self.lote, tipo_venda='ANALISE')
		alteracoes = [{
			'id': self.lote.id,
			'numero': self.lote.lote,
			'campos': {'area': {'atual': '200', 'novo': '999'}},
		}]
		self._post_confirmar(alteracoes)
		self.lote.refresh_from_db()
		self.assertEqual(self.lote.area, '200')

	def test_confirmar_json_invalido_redireciona_com_erro(self):
		response = self.client.post(self.url, {'alteracoes_json': 'invalido{'})
		self.assertRedirects(
			response,
			reverse('detalhe-empreendimento', args=[self.empr.uuid]),
			fetch_redirect_response=False,
		)


class ListarQuadrasAnaliseSemVendaTest(TestCase):
	"""Regressão: lote ANALISE sem RegisterVenda não pode crashar o link
	de proposta-rascunho do corretor (mesmo padrão do bug corrigido em analisa.html)."""

	def setUp(self):
		self.user = User.objects.create_user(
			username='corretor_test',
			password='pass123',
			tipo_usuario='CORRETOR',
		)
		self.client.force_login(self.user)
		self.empr = _make_empreendimento()
		self.quadra = _make_quadra(self.empr)
		self.lote = _make_lote(self.quadra, situacao=TypeLote.ANALISE)
		self.url = reverse('listar-quadras', args=[self.empr.uuid])

	def test_lote_analise_sem_venda_renderiza_sem_erro(self):
		response = self.client.get(self.url)
		self.assertEqual(response.status_code, 200)
		self.assertNotContains(response, 'documentos/proposta/rascunho/')

	def test_lote_analise_com_venda_gera_link_proposta_rascunho(self):
		venda = RegisterVenda.objects.create(lote=self.lote, tipo_venda='ANALISE')
		url_esperada = reverse('documentos:proposta-rascunho', args=[venda.uuid])
		response = self.client.get(self.url)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, url_esperada)